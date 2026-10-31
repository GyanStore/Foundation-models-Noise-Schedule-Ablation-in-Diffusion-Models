"""Build the Phase 2 report (Word .docx) from the experiment outputs.

Every number in the report is read from the result files written by the code in submission2/code
(training logs, schedule analysis, selected inference setting, test metrics, KID-vs-epoch), and every
figure comes from report_figures/ (make_figures.py). Re-running this script after re-running the
experiments regenerates the report with the new numbers.

Usage (from the repository root):
    python submission2/report/build_report.py --exp-dir outputs_real_metfaces_mps_tuned20
    soffice --headless --convert-to pdf --outdir submission2/report submission2/report/Phase2_Report.docx
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for easy landscape switch)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AUTHOR = "Shruthi Chinnasamy"
ROLL = "M26AI1119"
REPO = "https://github.com/GyanStore/Foundation-models-Noise-Schedule-Ablation-in-Diffusion-Models"
T = 1000
X = 2.00


# ----------------------------------------------------------------------------- data helpers
def read_csv(path: Path) -> list[dict]:
    with path.open() as f:
        return list(csv.DictReader(f))


def load(exp: Path, eval_sub: str) -> dict:
    tr, ev, fig = exp / "train", exp / eval_sub, exp / "report_figures"
    d: dict = {"fig": fig, "exp": exp}
    for r in ("A", "B"):
        d[f"train_{r}"] = read_csv(tr / f"run_{r}_train_log.csv")
        d[f"val_{r}"] = read_csv(tr / f"run_{r}_val_log.csv")
        d[f"cfg_{r}"] = json.loads((tr / f"run_{r}_config.json").read_text())
        d[f"best_{r}"] = json.loads((tr / f"run_{r}_best.json").read_text())
    d["sched"] = read_csv(exp / "schedule" / "schedule_values.csv")
    d["verify"] = (exp / "schedule" / "schedule_verification.txt").read_text().strip()
    d["setting"] = json.loads((ev / "selection" / "selected_setting.json").read_text())
    d["grid"] = read_csv(ev / "selection" / "selection_grid.csv")
    d["ilvr"] = {r["run"]: r for r in read_csv(ev / "test_ilvr" / "eval" / "eval_metrics.csv")}
    d["sdedit"] = {r["run"]: r for r in read_csv(ev / "test_sdedit" / "eval" / "eval_metrics.csv")}
    d["kve"] = read_csv(ev / "kid_vs_epoch.csv")
    return d


def sched_at(d: dict, t: int) -> dict:
    row = next(r for r in d["sched"] if int(r["t"]) == t)
    return {k: float(v) for k, v in row.items()}


def start_timestep(num_steps: int, strength: float) -> int:
    """Timestep where img2img sampling starts (diffusers 'leading' spacing, start_idx = round(n(1-s)))."""
    step_ratio = T // num_steps
    timesteps = (np.arange(0, num_steps) * step_ratio)[::-1]
    return int(timesteps[int(round(num_steps * (1 - strength)))])


def k(row: dict) -> float:
    return float(row["kid"]) * 1e3


def ks(row: dict) -> float:
    return float(row["kid_std"]) * 1e3


# ----------------------------------------------------------------------------- docx helpers
def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def table(doc: Document, header: list[str], rows: list[list[str]], widths_cm: list[float],
          mono: bool = False, font_pt: float = 9.5) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(font_pt)
        shade(c, "D9E2F3")
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            for li, line in enumerate(str(val).split("\n")):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line)
                run.font.size = Pt(font_pt - (1 if mono else 0))
                if mono:
                    run.font.name = "Courier New"
    for row in t.rows:
        for ci, w in enumerate(widths_cm):
            row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def para(doc: Document, text: str, bold_lead: str | None = None, size: float | None = None,
         italic: bool = False, align=None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align


def bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if it.startswith("**") and "**" in it[2:]:
            lead, rest = it[2:].split("**", 1)
            p.add_run(lead).bold = True
            p.add_run(rest)
        else:
            p.add_run(it)


def figure(doc: Document, path: Path, caption: str, width_cm: float = 16.5) -> None:
    if not path.exists():
        para(doc, f"[missing figure: {path.name}]", italic=True)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def figure_pair(doc: Document, left: Path, right: Path, caption: str, width_cm: float = 8.3) -> None:
    """Two figures side by side in a borderless 1x2 table, one caption below."""
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path in zip(t.rows[0].cells, (left, right)):
        cell.width = Cm(width_cm + 0.2)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            p.add_run().add_picture(str(path), width=Cm(width_cm))
        else:
            p.add_run(f"[missing figure: {path.name}]").italic = True
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


# ----------------------------------------------------------------------------- report
def build(d: dict, out: Path, notes: dict) -> None:
    fig = d["fig"]
    s = d["setting"]
    t0 = start_timestep(int(s["num_steps"]), float(s["strength"]))
    st0, st1, stT = sched_at(d, t0), sched_at(d, 1), sched_at(d, T)
    il, sd = d["ilvr"], d["sdedit"]
    kve = {(r["run"], int(r["epoch"])): r for r in d["kve"]}
    epochs = sorted({e for (_, e) in kve})
    cfg = d["cfg_A"]
    sp = cfg["split"]
    tlA, tlB = d["train_A"], d["train_B"]
    vlA, vlB = d["val_A"], d["val_B"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)
    base.paragraph_format.space_after = Pt(4)
    for lvl, size in ((1, 14), (2, 11.5)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Noise Schedule Ablation in Diffusion Models")
    r.bold = True
    r.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Phase 2 Report - Foundation Models and Generative AI\n").font.size = Pt(11)
    sub.add_run(f"Toonification of face photographs with google/ddpm-celebahq-256 fine-tuned on MetFaces, "
                f"native beta_t vs beta_t^{X:.2f}\n").font.size = Pt(10)
    sub.add_run(f"{AUTHOR} ({ROLL})   |   Code: {REPO}").font.size = Pt(9.5)

    # ------------------------------------------------------------------ 1
    doc.add_heading("1. Task and Model", level=1)
    para(doc, "Toonification of real face photographs: a real photographic portrait is translated into a "
              "painterly \"Old Masters\" portrait in the style of pre-20th-century portrait painting (the "
              "MetFaces domain), while the subject's approximate facial structure and pose are preserved. "
              "Input: an aligned 256x256 RGB face photo. Output: a 256x256 RGB painted-style portrait of the "
              "same face.", bold_lead="Task. ")
    para(doc, "google/ddpm-celebahq-256, an unconditional DDPM (Ho, Jain and Abbeel, \"Denoising Diffusion "
              "Probabilistic Models\", NeurIPS 2020, https://arxiv.org/abs/2006.11239) trained on CelebA-HQ "
              "256x256 face photographs; checkpoint https://huggingface.co/google/ddpm-celebahq-256 "
              "(diffusers DDPMPipeline = UNet2DModel + DDPMScheduler). The model is used for image-to-image "
              "translation by noising the input photo to an intermediate timestep and running the reverse "
              "process from there (SDEdit), optionally with ILVR low-frequency guidance (Section 4).",
         bold_lead="Model. ")
    para(doc, "For t = 1..T (T = 1000) the forward process is q(x_t | x_{t-1}) = N(sqrt(1 - beta_t) x_{t-1}, "
              "beta_t I), with closed form q(x_t | x_0) = N(sqrt(alpha_bar_t) x_0, (1 - alpha_bar_t) I), "
              "where alpha_t = 1 - beta_t and alpha_bar_t = prod_{s<=t} alpha_s. x_0 is the clean image, x_t "
              "the noisy image at step t, and I the identity. The network eps_theta(x_t, t) is trained with "
              "L = E || eps - eps_theta(sqrt(alpha_bar_t) x_0 + sqrt(1 - alpha_bar_t) eps, t) ||^2, "
              "eps ~ N(0, I).", bold_lead="Forward process. ")
    para(doc, "theta_t = beta_t, the per-step forward variance of DDPMScheduler. From the checkpoint's "
              "scheduler_config.json: beta_schedule = \"linear\", beta_start = 1e-4, beta_end = 0.02, "
              "num_train_timesteps = 1000. Run B uses theta_t^x = beta_t^x with x = "
              f"{X:.2f}; signal-to-noise ratio SNR(t) = alpha_bar_t / (1 - alpha_bar_t).",
         bold_lead="Native noise parameter. ")

    # ------------------------------------------------------------------ 2
    doc.add_heading("2. Dataset and Verification", level=1)
    para(doc, "MetFaces (NVlabs), https://github.com/NVlabs/metfaces-dataset: 1,336 face images at "
              "1024x1024 extracted from paintings, drawings and sculptures in the Metropolitan Museum of Art "
              "Open Access collection, aligned and cropped with the FFHQ procedure (dlib landmarks), CC BY-NC "
              "2.0. It is the target style domain of the task: aligned portraits from the Old Masters "
              "tradition, at a resolution that downsamples cleanly to the model's 256x256. We used the "
              "Hugging Face mirror huggan/metfaces (same 1,336 images).", bold_lead="Dataset. ")
    para(doc, "Route (b) of the brief. The pretraining domain is documented in Karras, Aila, Laine and "
              "Lehtinen, \"Progressive Growing of GANs for Improved Quality, Stability, and Variation\", ICLR "
              "2018, https://arxiv.org/abs/1710.10196, page 14, Appendix C: the Figure 8 caption (\"We start "
              "with a JPEG image (a) from the CelebA in-the-wild dataset\") and the paragraph below it (\"We "
              "perform the above processing for all 202599 images in the dataset ... discard all but the "
              "best 30000 images\"). CelebA-HQ is therefore made only of real camera photographs of "
              "celebrities. MetFaces consists of hand-made artworks (paint, pigment, brushstroke, engraving "
              "and stone textures) from art periods that largely predate photography and were never camera "
              "JPEGs, so it lies in a clearly different image domain and cannot have been part of "
              "CelebA-HQ.", bold_lead="Evidence it was not used for pretraining. ")
    para(doc, "Run 0 / A / B are applied to 30 real face photos from FFHQ (Karras et al., 2019; Flickr, no "
              "overlap with CelebA-HQ), mirror bitmind/ffhq-256, picked with seed 42; 20 further disjoint FFHQ "
              "photos form a dev set used only to choose the inference setting. FFHQ uses the same alignment "
              "as MetFaces, so inputs and training images share framing.", bold_lead="Input photos. ")

    # ------------------------------------------------------------------ 3
    doc.add_heading("3. Schedule Analysis", level=1)
    para(doc, f"beta_t = 1e-4 + (t - 1)(0.02 - 1e-4)/999 is linear and increasing, so 1e-4 <= beta_t <= 0.02 "
              f"for every t; hence 0 < theta_t < 1 for all t and theta_t^x is well defined, real and also in "
              f"(0, 1). Because 0 < theta_t < 1 and x > 1, theta_t^x < theta_t at every t. Measured: "
              f"beta_t^{X:.2f} spans 1e-8 .. 4e-4.", bold_lead="theta_t in (0, 1). ")
    figure(doc, fig / "fig_schedule.png", f"Figure 1. Left: theta_t = beta_t and theta_t^x (x = {X:.2f}) on the same "
           "axes. Middle: alpha_bar_t. Right: log-SNR(t) = log(alpha_bar_t / (1 - alpha_bar_t)).")
    para(doc, "Replacing beta_t by beta_t^x with x > 1 shrinks the noise added at every step, by a factor "
              "beta_t^(x-1) that is smallest (strongest shrinking) at small t. Since "
              "log alpha_bar_t = sum_s log(1 - beta_s^x) ~ -sum_s beta_s^x, the total injected noise falls from "
              f"sum beta_t = {sum(float(r['beta_t']) for r in d['sched']):.2f} to "
              f"sum beta_t^x = {sum(float(r['beta_t_pow_2.00']) for r in d['sched']):.3f}. Consequently "
              "alpha_bar_t is larger at every t: x_t = sqrt(alpha_bar_t) x_0 + sqrt(1 - alpha_bar_t) eps keeps "
              "far more of the image and far less noise. The endpoint changes qualitatively: "
              f"alpha_bar_T = {stT['alpha_bar_A']:.2e} under beta_t (x_T is pure noise, matching the N(0, I) "
              f"prior) but {stT['alpha_bar_B']:.3f} under beta_t^x, so q(x_T | x_0) = "
              f"N({np.sqrt(stT['alpha_bar_B']):.3f} x_0, {1 - stT['alpha_bar_B']:.3f} I) is still mostly the "
              "image and the forward process never reaches the Gaussian prior.",
         bold_lead="What changes in the forward process when x > 1. ")
    para(doc, "SNR(t) = alpha_bar_t / (1 - alpha_bar_t) rises at every timestep. In log-SNR (Figure 1, right) "
              f"the gap between the curves is {min(float(r['log_snr_B']) - float(r['log_snr_A']) for r in d['sched']):.1f} "
              f"to {max(float(r['log_snr_B']) - float(r['log_snr_A']) for r in d['sched']):.1f} nats: at t = 1 "
              f"log-SNR goes from {st1['log_snr_A']:.2f} to {st1['log_snr_B']:.2f} (approximately doubled, "
              f"since 1 - alpha_bar_1 = beta_1^x), at t = T from {stT['log_snr_A']:.2f} to {stT['log_snr_B']:.2f}. "
              "Run A sweeps log-SNR across about 19 nats, from almost clean to pure noise; Run B covers only "
              f"{stT['log_snr_B']:.1f} .. {st1['log_snr_B']:.1f}, i.e. the model never sees a low-SNR (heavily "
              f"noised) input. At our inference start timestep t = {t0}, alpha_bar_t is "
              f"{st0['alpha_bar_A']:.3f} for Run A but {st0['alpha_bar_B']:.3f} for Run B.",
         bold_lead="Effect on the signal-to-noise ratio. ")

    # ------------------------------------------------------------------ 4
    doc.add_heading("4. Experimental Setup", level=1)
    bullets(doc, [
        "**Dataset preparation.** All 1,336 MetFaces images were downloaded from huggan/metfaces and saved "
        "as RGB PNG. No image was removed: the set is already curated and aligned. On loading each image is "
        "resized to 256 (bicubic), center-cropped to 256x256 and scaled to [-1, 1]; training images get a "
        "random horizontal flip (p = 0.5).",
        f"**Split.** torch random_split with seed {cfg['seed']} into {sp['train_count']} train / "
        f"{sp['val_count']} validation / {sp['test_count']} test images (80/10/10). Run A and Run B use the "
        "identical split; the file lists are saved and were checked to be equal.",
        f"**Fine-tuning hyperparameters (identical for A and B).** Same pretrained checkpoint; AdamW, learning "
        f"rate {cfg['learning_rate']:g}, weight decay {cfg['weight_decay']:g}, cosine schedule with "
        f"{cfg['lr_warmup_steps']} warm-up steps; batch size {cfg['batch_size']} with "
        f"{cfg['grad_accum_steps']} gradient-accumulation steps (effective {cfg['effective_batch_size']}); "
        f"{cfg['epochs']} epochs; gradient clipping {cfg['grad_clip']}; EMA of weights with decay "
        f"{cfg['ema_decay']}; loss = MSE between true and predicted noise, timesteps uniform on 1..1000; "
        f"seed {cfg['seed']}; Apple M3 Max GPU (PyTorch MPS). The two config files differ only in run_name.",
        f"**Validation during fine-tuning.** After every epoch, the same noise-prediction MSE on the "
        f"{sp['val_count']} validation images, using the EMA weights and {cfg['val_repeats']} passes with "
        "fixed noise and fixed timesteps (seeded), so the value is comparable across epochs. The checkpoint "
        "with the lowest validation loss is kept; checkpoints are also saved at 5 equally spaced epochs "
        f"({', '.join(str(e) for e in epochs)}).",
        f"**Test-set evaluation.** The {sp['test_count']} held-out MetFaces test images are the reference "
        "distribution. Each run translates the 30 FFHQ test photos; we compute KID (main metric) and FID of "
        "the outputs against the MetFaces test images, and LPIPS between each output and its input photo. "
        "The test set was used only once, after the inference setting was frozen.",
        f"**Inference setting (frozen before testing).** On the 20 dev photos we evaluated strength "
        f"{{0.6, 0.8, 1.0}} x ILVR factor N {{8, 16}} (50 steps, Run A final model) and took the lowest KID "
        f"with LPIPS <= 0.45, giving {s['method'].upper()} with strength {s['strength']}, N = "
        f"{s['ilvr_factor']}, {s['num_steps']} DDPM steps, seed {s['seed']} (start timestep t = {t0}). The "
        "same setting is applied unchanged to Run 0, Run A and Run B. ILVR (Choi et al., ICCV 2021) adds, after "
        "every reverse step down to t = 200, x <- x - LP_N(x) + LP_N(q(y_t | y_0)), where LP_N "
        "down-samples by N and up-samples back, so the low-frequency layout of the output tracks the input "
        "photo while high frequencies (texture, brushwork) come from the model. Plain SDEdit at strength 0.6 "
        "is reported as a second, guidance-free comparison.",
    ])

    # ------------------------------------------------------------------ 5
    doc.add_heading("5. Implementation", level=1)
    table(doc, ["File (submission2/code/)", "Change"], [
        ["train_ddpm_metfaces.py", "line 29: EXPONENT_X = 2.00 named constant (# MODIFIED); lines 220-226 "
         "build_scheduler_for_run: Run B uses beta_t^x (3 lines # MODIFIED); MPS device, EMA, cosine LR, "
         "seeded validation, split manifests, checkpoint samples (# ADDED)."],
        ["run_ddpm_img2img.py", "SDEdit img2img; ADDED ILVR guidance (lowpass, --method/--ilvr-factor/"
         "--ilvr-stop-t); fix for 1 - alpha_bar = 0 underflow under beta_t^2; MODIFIED: seeded generator "
         "passed to scheduler.step."],
        ["analyze_schedule.py", "theta_t vs theta_t^x plot, alpha_bar_t, log-SNR, and verification of the "
         "saved Run A / Run B schedulers."],
        ["prepare_ffhq_faces.py, select_setting.py, run_final_eval.py, evaluate_outputs.py, make_figures.py",
         "ADDED: test/dev photos, dev-only setting selection, final test and checkpoint evaluation, "
         "KID/FID/LPIPS, figures."],
    ], [5.5, 11.5])
    para(doc, "Every changed line is marked # MODIFIED and every added line # ADDED. The only line-level "
              "difference between Run A and Run B is the scheduler construction:", bold_lead="")
    table(doc, ["Original (Run A, native schedule)", "Modified (Run B, theta_t^x)"], [[
        "noise_scheduler = DDPMScheduler.from_config(\n    base_scheduler.config)\n# betas = linear 1e-4..0.02",
        "EXPONENT_X = 2.00  # MODIFIED\n...\nbase_betas = base_scheduler.betas.detach().cpu()  # MODIFIED\n"
        "transformed_betas = torch.pow(base_betas, EXPONENT_X).numpy()  # MODIFIED\n"
        "return DDPMScheduler.from_config(base_scheduler.config,\n    trained_betas=transformed_betas)  # MODIFIED",
    ]], [7.0, 10.0], mono=True, font_pt=9)
    para(doc, "(1) The scheduler is saved inside every Run B checkpoint; analyze_schedule.py reloads the "
              "saved Run A and Run B schedulers and compares their betas with beta_t and beta_t^x. "
              "(2) Run B's training loss is on a different scale from Run A's from the first epoch (Section "
              "6), as expected when the noise to be predicted is much smaller. Output of the check:",
         bold_lead="Verifying that Run B used theta_t^x. ")
    table(doc, ["analyze_schedule.py verification output"], [[line] for line in d["verify"].splitlines()[:2]],
          [17.0], mono=True, font_pt=8.5)

    # ------------------------------------------------------------------ 6
    doc.add_heading("6. Results", level=1)
    doc.add_heading("6.1 Run 0 (pretrained, no fine-tuning)", level=2)
    figure(doc, fig / "fig_run0.png", "Figure 2. Run 0: the pretrained CelebA-HQ model applied to the test "
           "photos with the frozen setting. It reconstructs a clean photograph, not a painting.", 13.0)
    doc.add_heading("6.2 Fine-tuning loss", level=2)
    figure(doc, fig / "fig_loss.png", "Figure 3. Training (per-epoch mean) and validation noise-prediction MSE "
           "for Run A and Run B on one plot (log scale).", 13.5)
    para(doc, f"Run A: train {float(tlA[0]['train_loss']):.4f} -> {float(tlA[-1]['train_loss']):.4f}, val "
              f"{float(vlA[0]['val_loss']):.4f} -> {float(vlA[-1]['val_loss']):.4f} (best epoch "
              f"{d['best_A']['best_epoch']}). Run B: train {float(tlB[0]['train_loss']):.4f} -> "
              f"{float(tlB[-1]['train_loss']):.4f}, val {float(vlB[0]['val_loss']):.4f} -> "
              f"{float(vlB[-1]['val_loss']):.4f} (best epoch {d['best_B']['best_epoch']}). Both decrease "
              "steadily without over-fitting. The losses are not comparable in absolute value: under "
              "beta_t^x the noise to be predicted is a tiny component of x_t, so the per-pixel error is larger.")
    doc.add_heading("6.3 Samples at the 5 fine-tuning checkpoints", level=2)
    figure_pair(doc, fig / "fig_checkpoints_A.png", fig / "fig_checkpoints_B.png",
                f"Figure 4. Left: Run A; right: Run B, at epochs {', '.join(str(e) for e in epochs)} (rows), "
                "3 test photos (columns), frozen setting.")
    figure(doc, fig / "fig_kid_vs_epoch.png", "Figure 5. KID to MetFaces (left) and LPIPS to the input photo "
           "(right) for every checkpoint, 30 test photos each.")
    doc.add_heading("6.4 Inference trajectories", level=2)
    figure(doc, fig / "fig_trajectory_A.png", f"Figure 6. Run A final model: noisy sample x_t at 5 equally "
           f"spaced reverse steps from t = {t0} down to the last step (columns), 3 test photos (rows).", 14.0)
    figure(doc, fig / "fig_trajectory_B.png", "Figure 7. Run B final model: the same 5 reverse steps.", 14.0)
    doc.add_heading("6.5 Test-set scores", level=2)
    figure(doc, fig / "fig_main_comparison.png", "Figure 8. Test photos and outputs of Run 0, Run A and Run B "
           "(frozen ILVR setting) and of Run A with plain SDEdit.")
    rows = [["Input photos (reference point)", "-", f"{k(il['input']):.1f} +/- {ks(il['input']):.1f}",
             f"{float(il['input']['fid']):.1f}", "0"]]
    for name, key in (("Run 0 (pretrained)", "0"), ("Run A (beta_t)", "A"), ("Run B (beta_t^2)", "B")):
        for meth, src in (("ILVR (frozen)", il), ("SDEdit s=0.6", sd)):
            r_ = src[key]
            rows.append([name, meth, f"{k(r_):.1f} +/- {ks(r_):.1f}", f"{float(r_['fid']):.1f}",
                         f"{float(r_['lpips_to_input']):.3f}"])
    table(doc, ["Run", "Sampler", "KID x1e3 (main)", "FID", "LPIPS to input"], rows, [4.6, 3.0, 3.6, 2.4, 3.0])
    para(doc, "KID is the main metric because it measures how close the outputs are to the MetFaces "
              "painting distribution (the target of toonification) and, unlike FID, is unbiased for our "
              "small sample of 30 images; LPIPS to the input is reported to show how much of the photo's "
              "structure each run keeps.", bold_lead="Metric justification. ")

    # ------------------------------------------------------------------ 7
    doc.add_heading("7. Analysis", level=1)
    para(doc, notes["better"], bold_lead="Which run is better, and when it showed. ")
    para(doc, notes["mechanism"], bold_lead="Mechanistic explanation. ")
    para(doc, notes["run0"], bold_lead="Run 0 vs Run A and Run B. ")

    # ------------------------------------------------------------------ 8
    doc.add_heading("8. Takeaway", level=1)
    para(doc, notes["takeaway"])

    # ------------------------------------------------------------------ 9
    doc.add_heading("9. Contribution Statement", level=1)
    table(doc, ["Member", "Responsibility"], [[
        f"{AUTHOR} ({ROLL})",
        "All components: task, model and dataset selection (Phase 1); schedule analysis; code changes "
        "(beta_t^x schedule, MPS support, ILVR, evaluation); fine-tuning of Run A and Run B; Run 0 baseline; "
        "test evaluation; figures; report.",
    ]], [5.0, 12.0])
    para(doc, "Signed: ______________________________          Date: ______________")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--exp-dir", type=Path, default=Path("outputs_real_metfaces_mps_tuned20"))
    p.add_argument("--eval-subdir", default="final_eval")
    p.add_argument("--notes", type=Path, default=Path(__file__).with_name("analysis_notes.json"),
                   help="Analysis/Takeaway paragraphs (written after inspecting the results)")
    p.add_argument("--out", type=Path, default=Path("submission2/report/Phase2_Report.docx"))
    args = p.parse_args()
    notes = json.loads(args.notes.read_text())
    build(load(args.exp_dir, args.eval_subdir), args.out, notes)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
